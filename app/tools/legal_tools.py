"""
Legal Tools - Document reading and template operations for legal documents.

Provides tools for:
- Reading documents (.docx, .pdf, .txt)
- Filling templates with placeholders
"""
import re
from pathlib import Path
from langchain_core.tools import tool

from app.config import get_settings
from app.middleware.virtual_fs import get_vfs, resolve_path
from app.services import (
    session_store, 
    mark_document_as_read, 
    is_document_read,
    cache_document_content,
    get_cached_document_content
)


async def get_index_file_content(file_path: Path) -> str | None:
    """
    Check if an index file exists for the document (contains pre-extracted text).
    
    Looks in:
    1. Same directory: file.pdf.index.json
    2. indexedfiles folder: indexedfiles/file.pdf.index.json
    """
    settings = get_settings()
    file_name = file_path.name
    
    potential_paths = [
        file_path.parent / f"{file_name}.index.json",
        settings.indexed_files_dir / f"{file_name}.index.json",
        file_path.parent / "indexedfiles" / f"{file_name}.index.json",
    ]
    
    for index_path in potential_paths:
        try:
            if index_path.exists():
                import json
                content = index_path.read_text(encoding="utf-8")
                data = json.loads(content)
                if data.get("content"):
                    return f"[Read from Index]:\n{data['content']}"
        except Exception:
            continue
    
    return None


@tool
async def read_document(file_path: str, conversation_id: str = "default") -> str:
    """
    Read text from document files (.docx, .pdf, .txt, .md).
    Use this for reading specific legal documents.
    
    Supports virtual paths:
    - /project/document.pdf - Read from current project
    - /templates/contract.docx - Read from templates
    
    Args:
        file_path: Path to the file (virtual or absolute)
        conversation_id: Conversation ID for tracking
    """
    try:
        # Handle virtual paths
        if file_path.startswith("/"):
            try:
                path = resolve_path(file_path)
            except ValueError as e:
                return f"Error: {str(e)}"
        else:
            path = Path(file_path)
        
        file_name = path.name
        
        # Check if already read/cached in this session
        cached = await get_cached_document_content(session_store, conversation_id, file_name)
        if cached:
            return cached
        
        # Check for pre-indexed content
        indexed_content = await get_index_file_content(path)
        if indexed_content:
            await mark_document_as_read(session_store, conversation_id, file_name)
            await cache_document_content(session_store, conversation_id, file_name, indexed_content)
            return indexed_content
        
        # Fallback to raw parsing
        ext = path.suffix.lower()
        
        if ext == ".docx":
            from docx import Document
            doc = Document(str(path))
            content = "\n".join(para.text for para in doc.paragraphs if para.text.strip())
            result = content or "Empty DOCX file"
        
        elif ext == ".pdf":
            import fitz  # PyMuPDF
            doc = fitz.open(str(path))
            text_parts = []
            for page in doc:
                text_parts.append(page.get_text())
            content = "\n".join(text_parts)
            result = content or "Empty PDF file"
        
        else:
            # Default to text read
            result = path.read_text(encoding="utf-8", errors="ignore")
        
        # Mark as read and CACHE CONTENT
        await mark_document_as_read(session_store, conversation_id, file_name)
        await cache_document_content(session_store, conversation_id, file_name, result)
        
        return result
        
    except ImportError as e:
        return f"Missing dependency: {str(e)}. Install with pip."
    except Exception as e:
        return f"Error reading document: {str(e)}"


@tool
async def fill_template(template_content: str, replacements: str) -> str:
    """
    Replace placeholders like {{NAME}} in a template string with provided values.
    
    Args:
        template_content: The original template content (HTML/Text)
        replacements: JSON string of replacements, e.g. '{"CLIENT": "John Doe"}'
    """
    try:
        import json
        
        result = template_content
        
        # Handle JSON string input
        if isinstance(replacements, str):
            data = json.loads(replacements)
        else:
            data = replacements
        
        for key, value in data.items():
            # Replace {{KEY}} pattern
            pattern = r"\{\{" + re.escape(key) + r"\}\}"
            result = re.sub(pattern, str(value), result)
        
        return result
        
    except json.JSONDecodeError:
        return "Error: Replacements must be a valid JSON object."
    except Exception as e:
        return f"Error filling template: {str(e)}"


# Export all tools
LEGAL_TOOLS = [
    read_document,
    fill_template,
]
